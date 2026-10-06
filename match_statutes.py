#!/usr/bin/env python3
"""
Match LSC statutory citations to a directory of historical XML code files.

The script reads the long-format citation workbook created for the LSC project,
indexes XML files recursively, proposes a file match for each cited section,
and writes a new workbook plus a CSV review report. It never silently accepts
ties: ambiguous and low-confidence matches are retained for manual review.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import tempfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Sequence

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:
    raise SystemExit("This script requires openpyxl: pip install openpyxl") from exc


REQUIRED_COLUMNS = (
    "Jurisdiction",
    "Question",
    "Citation/Statute",
    "XML Title",
    "Chapter",
    "Section",
)

REVIEW_COLUMNS = (
    "XML File",
    "XML Year",
    "Available XML Years",
    "Historical XML Files",
    "Match Status",
    "Match Score",
    "Candidate XML Files",
    "Match Notes",
)

JURISDICTION_CODES = {
    "Alabama": "AL",
    "Alaska": "AK",
    "Arizona": "AZ",
    "Arkansas": "AR",
    "California": "CA",
    "Colorado": "CO",
    "Connecticut": "CT",
    "Delaware": "DE",
    "District of Columbia": "DC",
    "Florida": "FL",
    "Georgia": "GA",
    "Hawaii": "HI",
    "Idaho": "ID",
    "Illinois": "IL",
    "Indiana": "IN",
    "Iowa": "IA",
    "Kansas": "KS",
    "Kentucky": "KY",
    "Louisiana": "LA",
    "Maine": "ME",
    "Maryland": "MD",
    "Massachusetts": "MA",
    "Michigan": "MI",
    "Minnesota": "MN",
    "Mississippi": "MS",
    "Missouri": "MO",
    "Montana": "MT",
    "Nebraska": "NE",
    "Nevada": "NV",
    "New Hampshire": "NH",
    "New Jersey": "NJ",
    "New Mexico": "NM",
    "New York": "NY",
    "North Carolina": "NC",
    "North Dakota": "ND",
    "Ohio": "OH",
    "Oklahoma": "OK",
    "Oregon": "OR",
    "Pennsylvania": "PA",
    "Rhode Island": "RI",
    "South Carolina": "SC",
    "South Dakota": "SD",
    "Tennessee": "TN",
    "Texas": "TX",
    "Utah": "UT",
    "Vermont": "VT",
    "Virginia": "VA",
    "Washington": "WA",
    "West Virginia": "WV",
    "Wisconsin": "WI",
    "Wyoming": "WY",
}

SECTION_REFERENCE_RE = re.compile(
    r"(?:§{1,2}|\b(?:art(?:icle)?|rule)\.?\s+)"
    r"\s*([A-Z0-9][A-Z0-9.\-:]*"
    r"(?:\([A-Z0-9\-]+\))*)",
    re.IGNORECASE,
)

EXPLICIT_TITLE_RE = re.compile(
    r"\b(?:tit(?:le)?\.?|ch(?:apter)?\.?)\s*([A-Z0-9]+(?:[.\-:][A-Z0-9]+)*)",
    re.IGNORECASE,
)

SECTION_TAG_TERMS = (
    "section",
    "sectno",
    "sectionnumber",
    "sectionnum",
    "designator",
    "num",
    "number",
)

TITLE_TAG_TERMS = ("title", "code", "collection")
CHAPTER_TAG_TERMS = ("chapter", "article", "part")


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def path_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].lower()


def normalize_identifier(value: str) -> str:
    value = normalize_text(value).upper()
    value = value.replace("§", "")
    value = re.sub(r"^(?:SECTION|SEC\.?|ARTICLE|ART\.?|RULE)\s*", "", value)
    value = re.sub(r"\s+", "", value)
    value = re.sub(r"[^A-Z0-9.():\-]", "", value)
    return value.strip(".:-")


def strip_subsections(value: str) -> str:
    return re.sub(r"(?:\([A-Z0-9\-]+\))+$", "", value, flags=re.IGNORECASE)


def identifier_variants(value: str) -> set[str]:
    normalized = normalize_identifier(value)
    base = strip_subsections(normalized)
    variants = {v for v in (normalized, base) if v}

    for item in list(variants):
        variants.add(re.sub(r"[.\-:]", "", item))
        parts = re.split(r"([.\-:])", item)
        without_zeroes = []
        for part in parts:
            if re.fullmatch(r"\d+", part):
                without_zeroes.append(str(int(part)))
            else:
                without_zeroes.append(part)
        variants.add("".join(without_zeroes))
        variants.add(re.sub(r"[.\-:]", "", "".join(without_zeroes)))

    return {v for v in variants if len(v) >= 2}


def looks_like_legal_identifier(value: str) -> bool:
    value = normalize_identifier(value)
    if not value or len(value) > 50 or not any(ch.isdigit() for ch in value):
        return False
    return bool(re.fullmatch(r"[A-Z0-9.():\-]+", value))


def extract_section_references(citation: str) -> list[str]:
    references = []
    seen = set()
    for match in SECTION_REFERENCE_RE.finditer(citation or ""):
        reference = normalize_identifier(match.group(1))
        if reference and reference not in seen:
            references.append(reference)
            seen.add(reference)
    return references


def derive_title_hints(citation: str, references: Sequence[str]) -> set[str]:
    hints = {normalize_identifier(m.group(1)) for m in EXPLICIT_TITLE_RE.finditer(citation or "")}
    for reference in references:
        base = strip_subsections(reference)
        if ":" in base:
            hints.add(base.split(":", 1)[0])
        elif "-" in base:
            hints.add(base.split("-", 1)[0])
        elif "." in base:
            first = base.split(".", 1)[0]
            if len(first) <= 3:
                hints.add(first)
    return {normalize_identifier(h) for h in hints if h}


def derive_chapter_hint(section: str) -> str:
    section = strip_subsections(normalize_identifier(section))
    if ":" in section:
        remainder = section.split(":", 1)[1]
        return remainder.split("-", 1)[0]
    parts = section.split("-")
    if len(parts) >= 3:
        return parts[1]
    if "." in section:
        parts = section.split(".")
        if len(parts) >= 3:
            return parts[-2]
    return ""


def jurisdiction_aliases(jurisdiction: str) -> set[str]:
    aliases = {path_token(jurisdiction)}
    code = JURISDICTION_CODES.get(jurisdiction)
    if code:
        aliases.add(code.lower())
    aliases.add(path_token(jurisdiction.replace("District of Columbia", "Washington DC")))
    return {a for a in aliases if a}


def path_contains_alias(path: Path, aliases: set[str]) -> bool:
    components = {path_token(part) for part in path.parts}
    stem_tokens = {path_token(token) for token in re.split(r"[^A-Za-z0-9]+", path.stem)}
    tokens = components | stem_tokens
    return any(alias in tokens for alias in aliases)


@dataclass
class XmlRecord:
    path: Path
    relative_path: str
    path_compact: str
    year: int | None = None
    identifiers: set[str] = field(default_factory=set)
    title_values: set[str] = field(default_factory=set)
    chapter_values: set[str] = field(default_factory=set)
    parse_error: str = ""

    def contains_identifier(self, reference: str) -> bool:
        return bool(identifier_variants(reference) & self.identifiers)


@dataclass
class Candidate:
    record: XmlRecord
    score: int
    reasons: list[str]


@dataclass
class ReferenceMatch:
    reference: str
    status: str
    score: int
    matched: XmlRecord | None
    candidates: list[Candidate]
    selected_year: int | None


def add_identifier(target: set[str], value: object) -> None:
    text = normalize_text(value)
    if not looks_like_legal_identifier(text):
        return
    target.update(identifier_variants(text))


def extract_year(path: Path, root: Path) -> int | None:
    """Return the deepest four-digit year directory in a relative XML path."""
    relative = path.relative_to(root)
    years = []
    for part in relative.parts[:-1]:
        match = re.fullmatch(r"(?:19|20)\d{2}", part.strip())
        if match:
            years.append(int(match.group(0)))
    return years[-1] if years else None


def parse_xml_record(path: Path, root: Path) -> XmlRecord:
    relative = str(path.relative_to(root))
    record = XmlRecord(
        path=path,
        relative_path=relative,
        path_compact=path_token(relative),
        year=extract_year(path, root),
    )

    try:
        for _, element in ET.iterparse(path, events=("end",)):
            tag = local_name(element.tag)
            text = normalize_text(element.text)

            for key, value in element.attrib.items():
                key_name = local_name(key)
                if any(term in key_name for term in SECTION_TAG_TERMS):
                    add_identifier(record.identifiers, value)

            if any(term == tag or tag.endswith(term) for term in SECTION_TAG_TERMS):
                add_identifier(record.identifiers, text)
            if any(term == tag or tag.endswith(term) for term in TITLE_TAG_TERMS):
                if text and len(text) <= 150:
                    record.title_values.add(text)
            if any(term == tag or tag.endswith(term) for term in CHAPTER_TAG_TERMS):
                if text and len(text) <= 150:
                    record.chapter_values.add(text)

            element.clear()
    except (ET.ParseError, OSError) as exc:
        record.parse_error = str(exc)

    return record


def build_xml_index(xml_root: Path) -> list[XmlRecord]:
    paths = sorted(p for p in xml_root.rglob("*.xml") if p.is_file())
    if not paths:
        raise FileNotFoundError(f"No XML files found under {xml_root}")

    records = []
    for number, path in enumerate(paths, start=1):
        records.append(parse_xml_record(path, xml_root))
        if number % 250 == 0:
            print(f"Indexed {number:,}/{len(paths):,} XML files", file=sys.stderr)
    return records


def citation_tokens(citation: str) -> set[str]:
    ignored = {
        "code", "stat", "rev", "gen", "civil", "title", "tit", "section",
        "article", "art", "law", "laws", "ann", "compiled", "property",
    }
    return {
        token.lower()
        for token in re.findall(r"[A-Za-z]{2,}", citation or "")
        if token.lower() not in ignored
    }


def score_record(
    record: XmlRecord,
    jurisdiction: str,
    citation: str,
    reference: str,
    title_hints: set[str],
) -> Candidate:
    score = 0
    reasons = []
    aliases = jurisdiction_aliases(jurisdiction)

    if path_contains_alias(record.path, aliases):
        score += 30
        reasons.append("jurisdiction in path")

    exact_variants = identifier_variants(reference)
    if exact_variants & record.identifiers:
        score += 80
        reasons.append("section identifier found in XML structure")

    title_text = " ".join(record.title_values)
    title_compact = path_token(title_text)
    title_match = False
    for hint in title_hints:
        compact_hint = path_token(hint)
        if compact_hint and (
            compact_hint in record.path_compact
            or compact_hint in title_compact
        ):
            title_match = True
            break
    if title_match:
        score += 25
        reasons.append("title hint matched")

    overlap = 0
    path_words = set(re.findall(r"[a-z]{2,}", record.relative_path.lower()))
    title_words = set(re.findall(r"[a-z]{2,}", title_text.lower()))
    for token in citation_tokens(citation):
        if token in path_words or token in title_words:
            overlap += 1
    if overlap:
        score += min(15, overlap * 3)
        reasons.append(f"{overlap} citation term(s) matched")

    return Candidate(record=record, score=score, reasons=reasons)


def match_reference(
    records: Sequence[XmlRecord],
    jurisdiction: str,
    citation: str,
    reference: str,
    min_score: int,
    ambiguity_margin: int,
    max_candidates: int,
    target_year: int,
    year_policy: str,
) -> ReferenceMatch:
    title_hints = derive_title_hints(citation, [reference])
    jurisdiction_records = [
        record for record in records
        if path_contains_alias(record.path, jurisdiction_aliases(jurisdiction))
    ]
    pool = jurisdiction_records or list(records)

    selected_year = choose_snapshot_year(pool, target_year, year_policy)
    if selected_year is not None:
        pool = [record for record in pool if record.year == selected_year]
    elif year_policy == "exact" and any(record.year is not None for record in pool):
        return ReferenceMatch(reference, "NO_MATCH", 0, None, [], None)

    section_records = [record for record in pool if record.contains_identifier(reference)]
    if section_records:
        pool = section_records

    candidates = sorted(
        (
            score_record(record, jurisdiction, citation, reference, title_hints)
            for record in pool
        ),
        key=lambda candidate: (-candidate.score, candidate.record.relative_path.lower()),
    )[:max_candidates]

    if not candidates or candidates[0].score < min_score:
        score = candidates[0].score if candidates else 0
        return ReferenceMatch(reference, "NO_MATCH", score, None, candidates, selected_year)

    top = candidates[0]
    if len(candidates) > 1 and top.score - candidates[1].score < ambiguity_margin:
        return ReferenceMatch(reference, "AMBIGUOUS", top.score, None, candidates, selected_year)

    return ReferenceMatch(reference, "MATCHED", top.score, top.record, candidates, selected_year)


def choose_snapshot_year(
    records: Sequence[XmlRecord],
    target_year: int,
    year_policy: str,
) -> int | None:
    years = sorted({record.year for record in records if record.year is not None})
    if not years or year_policy == "all":
        return None
    if target_year in years:
        return target_year
    if year_policy == "exact":
        return None
    if year_policy == "nearest":
        return min(years, key=lambda year: (abs(year - target_year), year > target_year, year))

    prior = [year for year in years if year <= target_year]
    if prior:
        return max(prior)
    return min(years)


def historical_files_for_reference(
    records: Sequence[XmlRecord],
    jurisdiction: str,
    reference: str,
) -> list[XmlRecord]:
    aliases = jurisdiction_aliases(jurisdiction)
    matches = [
        record
        for record in records
        if path_contains_alias(record.path, aliases)
        and record.contains_identifier(reference)
    ]
    return sorted(
        matches,
        key=lambda record: (
            record.year is None,
            record.year if record.year is not None else 9999,
            record.relative_path.lower(),
        ),
    )


def unique_join(values: Iterable[str]) -> str:
    output = []
    seen = set()
    for value in values:
        value = normalize_text(value)
        if value and value not in seen:
            output.append(value)
            seen.add(value)
    return "; ".join(output)


def match_citation(
    records: Sequence[XmlRecord],
    jurisdiction: str,
    citation: str,
    min_score: int,
    ambiguity_margin: int,
    max_candidates: int,
    target_year: int,
    year_policy: str,
) -> dict[str, object]:
    references = extract_section_references(citation)
    if not normalize_text(citation):
        return {
            "status": "NO_CITATION",
            "score": 0,
            "xml_titles": "",
            "chapters": "",
            "sections": "",
            "xml_files": "",
            "xml_years": "",
            "available_years": "",
            "historical_files": "",
            "candidate_files": "",
            "notes": "The LSC citation field is blank.",
        }
    if not references:
        return {
            "status": "UNPARSED_CITATION",
            "score": 0,
            "xml_titles": "",
            "chapters": "",
            "sections": "",
            "xml_files": "",
            "xml_years": "",
            "available_years": "",
            "historical_files": "",
            "candidate_files": "",
            "notes": "No section, article, or rule identifier was parsed from the citation.",
        }

    matches = [
        match_reference(
            records,
            jurisdiction,
            citation,
            reference,
            min_score,
            ambiguity_margin,
            max_candidates,
            target_year,
            year_policy,
        )
        for reference in references
    ]

    matched = [match for match in matches if match.matched is not None]
    statuses = {match.status for match in matches}
    if statuses == {"MATCHED"}:
        status = "MATCHED"
    elif matched:
        status = "PARTIAL"
    elif "AMBIGUOUS" in statuses:
        status = "AMBIGUOUS"
    else:
        status = "NO_MATCH"

    matched_records = [match.matched for match in matched if match.matched]
    historical_records = [
        record
        for reference in references
        for record in historical_files_for_reference(records, jurisdiction, reference)
    ]
    candidate_files = unique_join(
        candidate.record.relative_path
        for match in matches
        for candidate in match.candidates
    )
    notes = unique_join(
        f"{match.reference}: {match.status}"
        + (
            f" ({', '.join(match.candidates[0].reasons)})"
            if match.candidates and match.candidates[0].reasons
            else ""
        )
        for match in matches
    )
    selected_year_note = unique_join(
        f"{match.reference}: snapshot {match.selected_year}"
        for match in matches
        if match.selected_year is not None
    )
    if selected_year_note:
        notes = unique_join([notes, selected_year_note])

    return {
        "status": status,
        "score": min((match.score for match in matches), default=0),
        "xml_titles": unique_join(record.path.stem for record in matched_records),
        "chapters": unique_join(derive_chapter_hint(reference) for reference in references),
        "sections": unique_join(references),
        "xml_files": unique_join(record.relative_path for record in matched_records),
        "xml_years": unique_join(str(record.year) for record in matched_records if record.year is not None),
        "available_years": unique_join(
            str(record.year) for record in historical_records if record.year is not None
        ),
        "historical_files": unique_join(
            f"{record.year if record.year is not None else 'unknown'}:{record.relative_path}"
            for record in historical_records
        ),
        "candidate_files": candidate_files,
        "notes": notes,
    }


def header_map(worksheet) -> dict[str, int]:
    return {
        normalize_text(cell.value): cell.column
        for cell in worksheet[1]
        if normalize_text(cell.value)
    }


def ensure_review_columns(worksheet, headers: dict[str, int]) -> dict[str, int]:
    source_header = worksheet.cell(1, max(headers.values()))
    for column_name in REVIEW_COLUMNS:
        if column_name in headers:
            continue
        column = worksheet.max_column + 1
        cell = worksheet.cell(1, column, column_name)
        if source_header.has_style:
            cell._style = copy(source_header._style)
        if source_header.number_format:
            cell.number_format = source_header.number_format
        headers[column_name] = column
    return headers


def write_match_to_row(
    worksheet,
    row_number: int,
    headers: dict[str, int],
    result: dict[str, object],
    overwrite_existing: bool,
) -> None:
    values = {
        "XML Title": result["xml_titles"],
        "Chapter": result["chapters"],
        "Section": result["sections"],
        "XML File": result["xml_files"],
        "XML Year": result["xml_years"],
        "Available XML Years": result["available_years"],
        "Historical XML Files": result["historical_files"],
        "Match Status": result["status"],
        "Match Score": result["score"],
        "Candidate XML Files": result["candidate_files"],
        "Match Notes": result["notes"],
    }
    for name, value in values.items():
        cell = worksheet.cell(row_number, headers[name])
        if name in REQUIRED_COLUMNS and not overwrite_existing and normalize_text(cell.value):
            continue
        cell.value = value


def process_workbook(args: argparse.Namespace) -> tuple[Path, Path, dict[str, int]]:
    workbook_path = Path(args.workbook).expanduser().resolve()
    xml_root = Path(args.xml_root).expanduser().resolve()
    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else workbook_path.with_name(f"{workbook_path.stem}_matched.xlsx")
    )
    report_path = (
        Path(args.report).expanduser().resolve()
        if args.report
        else output_path.with_suffix(".matches.csv")
    )

    records = build_xml_index(xml_root)
    workbook = load_workbook(workbook_path)
    worksheet = workbook[args.sheet] if args.sheet else workbook.active
    headers = header_map(worksheet)

    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Workbook is missing required column(s): {', '.join(missing)}")
    headers = ensure_review_columns(worksheet, headers)

    selected_jurisdictions = set(args.jurisdiction or [])
    report_rows = []
    summary = defaultdict(int)

    for row_number in range(2, worksheet.max_row + 1):
        jurisdiction = normalize_text(worksheet.cell(row_number, headers["Jurisdiction"]).value)
        question = normalize_text(worksheet.cell(row_number, headers["Question"]).value)
        citation = normalize_text(worksheet.cell(row_number, headers["Citation/Statute"]).value)
        if not jurisdiction:
            continue
        if selected_jurisdictions and jurisdiction not in selected_jurisdictions:
            continue

        result = match_citation(
            records,
            jurisdiction,
            citation,
            args.min_score,
            args.ambiguity_margin,
            args.max_candidates,
            args.target_year,
            args.year_policy,
        )
        summary[str(result["status"])] += 1
        write_match_to_row(
            worksheet,
            row_number,
            headers,
            result,
            args.overwrite_existing,
        )
        report_rows.append(
            {
                "Workbook Row": row_number,
                "Jurisdiction": jurisdiction,
                "Question": question,
                "Citation/Statute": citation,
                "Parsed Sections": result["sections"],
                "XML File": result["xml_files"],
                "XML Year": result["xml_years"],
                "Available XML Years": result["available_years"],
                "Historical XML Files": result["historical_files"],
                "Match Status": result["status"],
                "Match Score": result["score"],
                "Candidate XML Files": result["candidate_files"],
                "Match Notes": result["notes"],
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    with report_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(report_rows[0]) if report_rows else [
            "Workbook Row", "Jurisdiction", "Question", "Citation/Statute",
            "Parsed Sections", "XML File", "XML Year", "Available XML Years",
            "Historical XML Files", "Match Status", "Match Score",
            "Candidate XML Files", "Match Notes",
        ])
        writer.writeheader()
        writer.writerows(report_rows)

    return output_path, report_path, dict(summary)


def run_self_test() -> None:
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        xml_root = root / "xml"
        (xml_root / "Alabama" / "2021").mkdir(parents=True)
        (xml_root / "Alabama" / "2019").mkdir(parents=True)
        (xml_root / "Alaska" / "2020").mkdir(parents=True)

        (xml_root / "Alabama" / "2021" / "title_35.xml").write_text(
            "<code><title>Title 35</title><section><num>§ 35-9A-421</num>"
            "<heading>Noncompliance with rental agreement</heading></section></code>",
            encoding="utf-8",
        )
        (xml_root / "Alabama" / "2019" / "title_35.xml").write_text(
            "<code><title>Title 35</title><section><num>§ 35-9A-421</num>"
            "<heading>Noncompliance with rental agreement</heading></section></code>",
            encoding="utf-8",
        )
        (xml_root / "Alaska" / "2020" / "title_09.xml").write_text(
            "<code><title>Title 09</title><section number='09.45.100'>"
            "<heading>Notice to quit</heading></section></code>",
            encoding="utf-8",
        )

        records = build_xml_index(xml_root)
        alabama = match_citation(
            records,
            "Alabama",
            "Ala. Code § 35-9A-421. Noncompliance with rental agreement.",
            70,
            10,
            5,
            2021,
            "exact-or-prior",
        )
        alaska = match_citation(
            records,
            "Alaska",
            "Alaska Stat. § 09.45.100. Notice to quit.",
            70,
            10,
            5,
            2021,
            "exact-or-prior",
        )

        assert alabama["status"] == "MATCHED", alabama
        assert alabama["xml_files"] == "Alabama/2021/title_35.xml", alabama
        assert alabama["xml_years"] == "2021", alabama
        assert alabama["available_years"] == "2019; 2021", alabama
        assert alabama["sections"] == "35-9A-421", alabama
        assert alaska["status"] == "MATCHED", alaska
        assert alaska["xml_files"] == "Alaska/2020/title_09.xml", alaska
        assert alaska["xml_years"] == "2020", alaska
        assert alaska["sections"] == "09.45.100", alaska

    print("Self-test passed.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Match LSC citations to historical state-statute XML files.",
    )
    parser.add_argument("--workbook", help="Input LSC citation workbook (.xlsx).")
    parser.add_argument("--xml-root", help="Root directory containing XML files.")
    parser.add_argument("--output", help="Output workbook path.")
    parser.add_argument("--report", help="Optional CSV review-report path.")
    parser.add_argument("--sheet", help="Worksheet name. Defaults to the active sheet.")
    parser.add_argument(
        "--jurisdiction",
        action="append",
        help="Process one jurisdiction only. Repeat the flag for multiple jurisdictions.",
    )
    parser.add_argument(
        "--target-year",
        type=int,
        default=2021,
        help="Snapshot year used to anchor the LSC citations. Default: 2021.",
    )
    parser.add_argument(
        "--year-policy",
        choices=("exact", "exact-or-prior", "nearest", "all"),
        default="exact-or-prior",
        help=(
            "How to select a snapshot when the target year is missing. "
            "Default: exact-or-prior."
        ),
    )
    parser.add_argument(
        "--min-score",
        type=int,
        default=70,
        help="Minimum score required for a match. Default: 70.",
    )
    parser.add_argument(
        "--ambiguity-margin",
        type=int,
        default=10,
        help="Minimum lead over the second candidate. Default: 10.",
    )
    parser.add_argument(
        "--max-candidates",
        type=int,
        default=5,
        help="Maximum candidates retained for review. Default: 5.",
    )
    parser.add_argument(
        "--overwrite-existing",
        action="store_true",
        help="Replace existing XML Title, Chapter, and Section values.",
    )
    parser.add_argument(
        "--self-test",
        action="store_true",
        help="Run a small built-in matching test and exit.",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    if args.self_test:
        run_self_test()
        return 0
    if not args.workbook or not args.xml_root:
        parser.error("--workbook and --xml-root are required unless --self-test is used")

    output, report, summary = process_workbook(args)
    print(f"Workbook: {output}")
    print(f"Review report: {report}")
    print("Match summary:")
    for status, count in sorted(summary.items()):
        print(f"  {status}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())